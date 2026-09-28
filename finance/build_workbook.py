"""Builds the Pod 2 Financial & NFR Workbook (all calculations are live Excel formulas)."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(sys.argv[1])
M = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))            # MockLLM run (guardrail demo)
RUN_TS = M["generated_at"]
# Optional real-model run (AUTOHEAL_LLM=litellm): latency, cost and runtime metrics come from it.
REAL = json.loads(Path(sys.argv[3]).read_text(encoding="utf-8")) if len(sys.argv) > 3 else None
DEMO_INCIDENTS = 6

ARIAL = "Arial"
F_IN = Font(name=ARIAL, size=10, color="0000FF")
F_CALC = Font(name=ARIAL, size=10, color="000000")
F_CALC_B = Font(name=ARIAL, size=10, color="000000", bold=True)
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_LINK_B = Font(name=ARIAL, size=10, color="008000", bold=True)
F_HDR = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True, color="1F3864")
F_SEC = Font(name=ARIAL, size=11, bold=True, color="1F3864")
F_TXT = Font(name=ARIAL, size=10)
F_TXT_B = Font(name=ARIAL, size=10, bold=True)
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
F_FORMULA = Font(name="Consolas", size=11, bold=True, color="1F3864")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_TOT = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))
WRAP = Alignment(wrap_text=True, vertical="top")

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
USDK = '$#,##0.000;($#,##0.000);"-"'
PCT = '0.0%;(0.0%);"-"'
INT = '#,##0;(#,##0);"-"'
DEC1 = '#,##0.0;(#,##0.0);"-"'
DEC2 = '0.00;(0.00);"-"'
MULT = '0.0"x"'

R: dict[str, str] = {}          # named inputs → absolute references

wb = Workbook()
readme = wb.active
readme.title = "README"


def put(ws, ref, value, font=F_CALC, fmt=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font, c.fill = F_HDR, FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")


def section(ws, row, title, ncols):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = FILL_SEC
    put(ws, f"A{row}", title, F_SEC, fill=FILL_SEC)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# =====================================================================================
# Assumptions
# =====================================================================================
A = wb.create_sheet("Assumptions")
widths(A, {"A": 52, "B": 15, "C": 15, "D": 15, "E": 15, "F": 16, "G": 70})
put(A, "A1", "Assumptions & Inputs: Pod 2 AIOps Auto-Healing Platform", F_TITLE)
put(A, "A2", "Blue text = input you can change · Yellow fill = key lever · Black = formula · Green = link from "
             "another sheet. Every hardcoded number has its source or rationale in the last column.", F_NOTE)

section(A, 4, "1. Scenario selector", 7)
put(A, "A5", "Active scenario (choose Conservative / Base / Optimistic)", F_TXT_B)
put(A, "B5", "Base", F_IN, fill=FILL_KEY)
put(A, "G5", "Drives column E of the lever table below; every sheet recalculates.", F_NOTE)
dv = DataValidation(type="list", formula1='"Conservative,Base,Optimistic"', allow_blank=False)
A.add_data_validation(dv)
dv.add("B5")
R["scenario"] = "Assumptions!$B$5"

section(A, 7, "2. Key levers by scenario (these drive the ROI)", 7)
header(A, 8, ["Parameter", "Conservative", "Base", "Optimistic", "Active", "Unit", "Source / rationale"])
LEVERS = [
    ("auto_share", "Incidents fully auto-remediated (READ / NON_DESTRUCTIVE, C ≥ 0.85)", 0.20, 0.30, 0.40, "% of incidents", PCT,
     "Pod assumption. Prototype demo: 2 of 6 incidents auto-remediated (stale-cache ×2)."),
    ("hitl_share", "Incidents HITL-assisted (DESTRUCTIVE, SRE-approved)", 0.25, 0.30, 0.35, "% of incidents", PCT,
     "Pod assumption. Prototype demo: 3 of 6 incidents approved via the HITL gate."),
    ("mttr_red", "MTTR reduction on covered incidents", 0.30, 0.40, 0.55, "%", PCT,
     "Pod assumption: triage + RCA + plan drafted in minutes instead of manual diagnosis; validate in pilot."),
    ("downtime_cost", "Business cost of one P1 hour", 5000, 8000, 15000, "$ / hour", USD,
     "Pod assumption (revenue loss + SLA credits). CloudScale Finance to confirm; most sensitive input."),
    ("attribution", "Share of MTTR gain attributed to this platform", 0.40, 0.50, 0.60, "%", PCT,
     "Conservatism factor: part of any MTTR gain comes from other process changes."),
    ("toil_red", "Responder toil reduction on covered incidents", 0.30, 0.40, 0.55, "%", PCT,
     "Pod assumption: fewer engineers paged, diagnostics pre-collected by the Triage agent."),
    ("alert_auto", "Human-reviewed alerts auto-triaged / de-duplicated", 0.45, 0.60, 0.75, "%", PCT,
     "Pod assumption: alert-storm de-duplication via classification + semantic cache (ADR 03)."),
    ("realise", "Soft-savings realisation (freed hours turned into value)", 0.35, 0.50, 0.65, "%", PCT,
     "Freed engineer time is capacity, not cash: only this share is counted in toil and alert-triage benefits."),
]
lever_rows = {}
for i, (key, label, c, b, o, unit, fmt, note) in enumerate(LEVERS):
    r = 9 + i
    lever_rows[key] = r
    put(A, f"A{r}", label, F_TXT)
    for col, v in zip("BCD", (c, b, o)):
        put(A, f"{col}{r}", v, F_IN, fmt, FILL_KEY)
    put(A, f"E{r}", f"=INDEX($B{r}:$D{r},MATCH($B$5,$B$8:$D$8,0))", F_CALC_B, fmt)
    put(A, f"F{r}", unit, F_TXT)
    put(A, f"G{r}", note, F_NOTE, align=WRAP)
    R[key] = f"Assumptions!$E${r}"

row = 9 + len(LEVERS) + 1


def single_section(title, items):
    global row
    section(A, row, title, 7)
    row += 1
    header(A, row, ["Parameter", "Value", "Unit", "", "", "", "Source / rationale"])
    row += 1
    for key, label, value, unit, fmt, note in items:
        put(A, f"A{row}", label, F_TXT)
        put(A, f"B{row}", value, F_IN, fmt)
        put(A, f"C{row}", unit, F_TXT)
        put(A, f"G{row}", note, F_NOTE, align=WRAP)
        R[key] = f"Assumptions!$B${row}"
        row += 1
    row += 1


single_section("3. Volume & operations (Year 1 baseline)", [
    ("alerts_month", "Monitoring alerts per month (Prometheus / Datadog / CloudWatch)", 150000, "alerts / month", INT,
     "Pod assumption for a global multi-cloud estate (AWS + Azure)."),
    ("incidents_month", "Actionable incidents per month (all severities)", 600, "incidents / month", INT,
     "Pod assumption: incidents that need diagnosis and a remediation decision."),
    ("p1p2_share", "Share of incidents that are P1 / P2", 0.10, "%", PCT, "Pod assumption; brief targets P1/P2 MTTR."),
    ("p1_share", "P1 share within P1 / P2 incidents", 0.25, "%", PCT, "Pod assumption."),
    ("growth", "Annual growth in alert and incident volume", 0.15, "% / year", PCT, "Pod assumption (estate growth)."),
    ("alerts_human", "Share of alerts that need a human look today", 0.15, "%", PCT,
     "Pod assumption; the rest are auto-closed or ignored."),
    ("min_alert", "Minutes per human-reviewed alert", 2, "minutes", DEC1, "Pod assumption."),
    ("mttr_p1", "Baseline MTTR, P1", 3.0, "hours", DEC1, "Pod assumption (current state, manual diagnosis)."),
    ("mttr_p2", "Baseline MTTR, P2", 6.0, "hours", DEC1, "Pod assumption (current state)."),
    ("p2_factor", "P2 business impact relative to a P1 hour", 0.10, "x of P1 cost", PCT,
     "Pod assumption: P2 = degraded service, not outage."),
    ("eng_hours", "Responder hours per incident (all engineers involved)", 3.0, "hours", DEC1, "Pod assumption."),
    ("sre_rate", "Loaded SRE cost per hour", 95, "$ / hour", USD, "Equals $190k loaded FTE / ~2,000 h. Pod assumption."),
    ("hitl_min", "SRE review time per HITL approval", 10, "minutes", DEC1,
     "Pod assumption: approval card carries RCA, diff, blast radius, rollback."),
    ("adopt1", "Adoption ramp, Year 1 (share of estate onboarded)", 0.50, "%", PCT, "Phased rollout: pilot clusters first."),
    ("adopt2", "Adoption ramp, Year 2", 0.85, "%", PCT, "Pod assumption."),
    ("adopt3", "Adoption ramp, Year 3", 1.00, "%", PCT, "Pod assumption."),
    ("discount", "Discount rate for NPV", 0.10, "% / year", PCT, "Pod assumption (typical corporate hurdle rate)."),
    ("peak", "Peak-to-average traffic factor (alert storms)", 6, "x", MULT, "Pod assumption for capacity planning."),
    ("esc_rate", "Tier-1 → Tier-2 escalation rate", 0.10, "% of Tier-1 calls", PCT,
     "ADR 03: Tier-1 answers with confidence < 0.7 are retried on Tier 2."),
    ("infra_growth", "Annual growth in infrastructure cost", 0.10, "% / year", PCT, "Pod assumption (scales with estate)."),
    ("tpm_quota", "Provisioned LLM throughput quota", 20000, "tokens / sec", INT,
     "Pod assumption (provider rate limit to request); checked in NFR matrix."),
])

single_section("4. LLM pricing ($ per 1M tokens)", [
    ("t1_in", "Tier 1 (small/fast, e.g. GPT-4o-mini): input", 0.15, "$ / 1M tokens", USD2,
     "Provider list price used for this model; verify current pricing before presenting."),
    ("t1_out", "Tier 1: output", 0.60, "$ / 1M tokens", USD2, "As above."),
    ("t2_in", "Tier 2 (large/reasoning, e.g. GPT-4o): input", 2.50, "$ / 1M tokens", USD2, "As above."),
    ("t2_out", "Tier 2: output", 10.00, "$ / 1M tokens", USD2, "As above."),
    ("emb", "Embeddings for semantic cache (e.g. text-embedding-3-small)", 0.02, "$ / 1M tokens", USD2, "As above."),
])

# ---- Compute_infra components
section(A, row, "5. Compute_infra components (monthly, Year 1)", 7)
row += 1
header(A, row, ["Component", "Qty", "Unit cost ($)", "Hours / month", "Monthly cost ($)", "Basis", "Source / rationale"])
row += 1
INFRA = [
    ("Amazon EKS control plane (agent runtime cluster)", 1, 0.10, 730, "per hour", "AWS EKS list price $0.10/cluster-hour."),
    ("EKS worker nodes (m6i.xlarge): orchestrator, MCP server, OPA, LiteLLM", 4, 0.192, 730, "per hour",
     "AWS on-demand list price (us-east-1, approx.); verify at aws.amazon.com/pricing."),
    ("ElastiCache Redis (r6g.large, HA pair): semantic cache + checkpointer", 2, 0.206, 730, "per hour",
     "AWS on-demand list price (approx.)."),
    ("RDS PostgreSQL + pgvector (db.r6g.large, Multi-AZ)", 2, 0.26, 730, "per hour", "AWS on-demand list price (approx.)."),
    ("Azure AKS footprint for MCP adapters (2 × D4s v5)", 2, 0.192, 730, "per hour", "Azure pay-as-you-go list price (approx.)."),
    ("Observability stack (managed Grafana / Tempo / Prometheus)", 1, 600, 1, "per month", "Pod estimate."),
    ("Audit storage (S3 Object Lock) + SIEM ingest", 1, 450, 1, "per month", "Pod estimate."),
    ("Networking: NAT gateways, cross-cloud egress", 1, 250, 1, "per month", "Pod estimate."),
]
infra_first = row
for label, qty, unit_cost, hours, basis, note in INFRA:
    put(A, f"A{row}", label, F_TXT)
    put(A, f"B{row}", qty, F_IN, INT)
    put(A, f"C{row}", unit_cost, F_IN, '$#,##0.000' if unit_cost < 10 else USD)
    put(A, f"D{row}", hours, F_IN, INT)
    put(A, f"E{row}", f"=B{row}*C{row}*D{row}", F_CALC, USD)
    put(A, f"F{row}", basis, F_TXT)
    put(A, f"G{row}", note, F_NOTE, align=WRAP)
    row += 1
put(A, f"A{row}", "Compute_infra: total per month", F_TXT_B, fill=FILL_TOT)
put(A, f"E{row}", f"=SUM(E{infra_first}:E{row - 1})", F_CALC_B, USD, FILL_TOT, border=TOP)
R["infra_month"] = f"Assumptions!$E${row}"
row += 2

single_section("6. Ops_overhead inputs", [
    ("platform_fte", "Platform operations engineers (run, upgrade, on-call for the platform)", 1.5, "FTE", DEC1,
     "Pod assumption."),
    ("eval_fte", "Model evaluation + guardrail / OPA policy upkeep", 0.5, "FTE", DEC1,
     "ADR 02/03: owner for deny-lists, policies, weekly Tier-1 vs Tier-2 evaluation."),
    ("fte_cost", "Loaded cost per FTE", 190000, "$ / FTE / year", USD, "Pod assumption."),
    ("salary_esc", "Annual salary escalation", 0.03, "% / year", PCT, "Pod assumption."),
    ("tooling", "LLM observability & eval tooling licences (e.g. LangSmith)", 6000, "$ / year", USD, "Pod estimate."),
    ("build_eng", "Build & integration team (one-time, Year 1)", 4, "engineers", INT, "Pod assumption."),
    ("build_months", "Build & integration duration", 4, "months", INT, "Pod assumption."),
    ("sec_review", "Security review, threat model & pen test (one-time)", 40000, "$", USD, "Pod estimate."),
])

# =====================================================================================
# Token Economics
# =====================================================================================
T = wb.create_sheet("Token Economics")
widths(T, {"A": 44, "B": 10, "C": 9, "D": 10, "E": 10, "F": 14, "G": 9, "H": 10, "I": 10, "J": 9, "K": 13,
           "L": 10, "M": 10, "N": 10, "O": 9, "P": 13, "Q": 13, "R": 14, "S": 14, "T": 58})
put(T, "A1", "Token Economics & Model Routing", F_TITLE)
put(T, "A2", "Token_cost = N_requests × ((T_in × P_in) + (T_out × P_out)) × (1 − Cache_ratio)", F_FORMULA)
put(T, "A3", "Formula as given in the Capstone brief, §4 Artifact 4. Applied per task row; prices are per 1M tokens, "
             "so each row divides by 1,000,000. Optimized Tier-1 prices include the Tier-2 escalation surcharge.", F_NOTE)
put(T, "A5", "Alerts per month", F_TXT)
put(T, "B5", f"={R['alerts_month']}", F_LINK, INT)
put(T, "A6", "Incidents per month", F_TXT)
put(T, "B6", f"={R['incidents_month']}", F_LINK, INT)
put(T, "A7", "Tier-1 → Tier-2 escalation rate", F_TXT)
put(T, "B7", f"={R['esc_rate']}", F_LINK, PCT)

for rng, lab in (("A8:F8", "Workload profile"), ("G8:K8", "BASELINE: all Tier 2, no cache"),
                 ("L8:Q8", "OPTIMIZED: tiered routing + semantic cache (ADR 03)")):
    T.merge_cells(rng)
    put(T, rng.split(":")[0], lab, F_HDR, fill=FILL_HDR, align=Alignment(horizontal="center"))
header(T, 9, ["Task (agent step)", "Unit of work", "Calls per unit", "T_in (tokens / call)", "T_out (tokens / call)",
              "N_requests / month", "Tier", "P_in ($/1M)", "P_out ($/1M)", "Cache_ratio", "Token_cost ($/mo)",
              "Tier", "P_in eff. ($/1M)", "P_out eff. ($/1M)", "Cache_ratio", "Token_cost ($/mo)", "Savings ($/mo)",
              "Raw tokens / mo", "Billed tokens / mo (opt.)", "Routing rationale"])
T.row_dimensions[9].height = 42
TASKS = [
    ("Alert classification (classify)", "alert", 1, 700, 80, "Tier 1", 0.55,
     "Short, repetitive prompt; cached per alert rule (alert storms)."),
    ("Alert de-duplication & correlation (Supervisor)", "alert", 1, 1200, 120, "Tier 1", 0.40,
     "Pattern matching over recent alerts; Tier 1 is sufficient."),
    ("Input guardrail L2 (LlamaGuard-class check)", "incident", 1, 7000, 20, "Tier 1", 0.00,
     "Runs once per untrusted input, not per agent hop (ADR 03)."),
    ("Log summarisation (Triage: summarize_logs)", "incident", 4, 8000, 600, "Tier 1", 0.35,
     "High volume, low reasoning; cached on log fingerprint."),
    ("Root-cause analysis (Triage: rca, tool loop)", "incident", 6, 12000, 800, "Tier 2", 0.30,
     "Hard multi-signal reasoning stays on Tier 2; repeat incidents hit cache."),
    ("Remediation plan + rollback (Planner)", "incident", 1, 5000, 1500, "Tier 2", 0.00,
     "Never cached, never downgraded: this step writes production changes."),
    ("Supervisor routing & budget checks", "incident", 4, 2000, 150, "Tier 1", 0.00, "Simple routing decisions."),
    ("Post-fix verification (Verifier)", "incident", 2, 1500, 200, "Tier 1", 0.00, "Health-signal interpretation."),
]
FIRST = 10
for i, (task, unit, calls, tin, tout, tier, cache, why) in enumerate(TASKS):
    r = FIRST + i
    put(T, f"A{r}", task, F_TXT)
    put(T, f"B{r}", unit, F_IN)
    put(T, f"C{r}", calls, F_IN, INT)
    put(T, f"D{r}", tin, F_IN, INT)
    put(T, f"E{r}", tout, F_IN, INT)
    put(T, f"F{r}", f'=IF(B{r}="alert",$B$5,$B$6)*C{r}', F_CALC, INT)
    put(T, f"G{r}", "Tier 2", F_IN)
    put(T, f"J{r}", 0, F_IN, PCT)
    put(T, f"L{r}", tier, F_IN, fill=FILL_KEY)
    put(T, f"O{r}", cache, F_IN, PCT, FILL_KEY)
    put(T, f"T{r}", why, F_NOTE)
EMB = FIRST + len(TASKS)
put(T, f"A{EMB}", "Semantic-cache embedding lookups", F_TXT)
put(T, f"B{EMB}", "lookup", F_IN)
put(T, f"C{EMB}", "-", F_TXT)
put(T, f"D{EMB}", 1000, F_IN, INT)
put(T, f"E{EMB}", 0, F_IN, INT)
put(T, f"F{EMB}", f'=SUMIFS(F{FIRST}:F{EMB - 1},O{FIRST}:O{EMB - 1},">0")', F_CALC, INT)
put(T, f"G{EMB}", "n/a", F_IN)
put(T, f"J{EMB}", 0, F_IN, PCT)
put(T, f"L{EMB}", "Embedding", F_IN)
put(T, f"O{EMB}", 0, F_IN, PCT)
put(T, f"T{EMB}", "One embedding per cacheable call (only exists in the optimized design).", F_NOTE)
for r in range(FIRST, EMB + 1):
    put(T, f"H{r}", f'=IF(G{r}="Tier 2",{R["t2_in"]},IF(G{r}="Tier 1",{R["t1_in"]},0))', F_CALC, USD2)
    put(T, f"I{r}", f'=IF(G{r}="Tier 2",{R["t2_out"]},IF(G{r}="Tier 1",{R["t1_out"]},0))', F_CALC, USD2)
    put(T, f"K{r}", f"=F{r}*((D{r}*H{r})+(E{r}*I{r}))/1000000*(1-J{r})", F_CALC, USD2)
    put(T, f"M{r}", f'=IF(L{r}="Tier 1",{R["t1_in"]}+$B$7*{R["t2_in"]},IF(L{r}="Tier 2",{R["t2_in"]},'
                    f'IF(L{r}="Embedding",{R["emb"]},0)))', F_CALC, USDK)
    put(T, f"N{r}", f'=IF(L{r}="Tier 1",{R["t1_out"]}+$B$7*{R["t2_out"]},IF(L{r}="Tier 2",{R["t2_out"]},0))',
        F_CALC, USDK)
    put(T, f"P{r}", f"=F{r}*((D{r}*M{r})+(E{r}*N{r}))/1000000*(1-O{r})", F_CALC, USD2)
    put(T, f"Q{r}", f"=K{r}-P{r}", F_CALC, USD2)
    llm = r < EMB
    put(T, f"R{r}", f"=F{r}*(D{r}+E{r})" if llm else 0, F_CALC, INT)
    put(T, f"S{r}", f"=R{r}*(1-O{r})" if llm else 0, F_CALC, INT)
TOT = EMB + 1
put(T, f"A{TOT}", "Total per month", F_TXT_B, fill=FILL_TOT)
for col in "KPQRS":
    put(T, f"{col}{TOT}", f"=SUM({col}{FIRST}:{col}{EMB})", F_CALC_B, USD2 if col in "KPQ" else INT, FILL_TOT, border=TOP)

s = TOT + 2
section(T, s, "Summary (monthly, Year 1 volumes)", 6)
SUMMARY = [
    ("tok_base", "Token_cost: baseline", f"=K{TOT}", USD),
    ("tok_opt", "Token_cost: optimized", f"=P{TOT}", USD),
    ("tok_sav", "Savings per month", f"=Q{TOT}", USD),
    ("tok_sav_pct", "Token optimisation savings %", f"=IF(K{TOT}=0,0,Q{TOT}/K{TOT})", PCT),
    ("cr_blend", "Blended cache ratio, CR (optimized)", f"=IF(R{TOT}=0,0,1-S{TOT}/R{TOT})", PCT),
    ("tier2_share", "Tier 2 share of optimized cost", f'=IF(P{TOT}=0,0,SUMIFS(P{FIRST}:P{EMB},L{FIRST}:L{EMB},"Tier 2")/P{TOT})', PCT),
    ("tps_avg", "Tokens / sec: average (billed, optimized)", f"=S{TOT}/(30*24*3600)", DEC1),
    ("tps_peak", "Tokens / sec: peak (× peak factor)", None, DEC1),
    ("cost_per_inc", "LLM cost per incident (optimized, incl. alert share)", f"=IF($B$6=0,0,P{TOT}/$B$6)", USD2),
]
for i, (key, label, formula, fmt) in enumerate(SUMMARY):
    r = s + 1 + i
    put(T, f"A{r}", label, F_TXT)
    if key == "tps_peak":
        formula = f"=B{r - 1}*{R['peak']}"
    put(T, f"B{r}", formula, F_CALC_B, fmt)
    R[key] = f"'Token Economics'!$B${r}"
T.freeze_panes = "B10"

# =====================================================================================
# TCO & ROI
# =====================================================================================
X = wb.create_sheet("TCO & ROI")
widths(X, {"A": 58, "B": 15, "C": 15, "D": 15, "E": 16, "F": 80})
put(X, "A1", "3-Year TCO & ROI", F_TITLE)
put(X, "A2", "TCO = Compute_infra + Token_cost + Ops_overhead", F_FORMULA)
put(X, "A3", "Formula as given in the Capstone brief, §4 Artifact 4. All $ figures are nominal; NPV discounts net benefit.", F_NOTE)
header(X, 4, ["Line item", "Year 1", "Year 2", "Year 3", "3-Year Total", "Formula / note"])
rows: dict[str, int] = {}
cur = 5
YC = "BCD"


def line(key, label, fn, fmt=USD, total=True, bold=False, note="", font=None, fill=None, border=None):
    """fn(col, year_index_ref) -> formula string for that year column."""
    global cur
    r = cur
    rows[key] = r
    put(X, f"A{r}", label, F_TXT_B if bold else F_TXT, fill=fill)
    for col in YC:
        put(X, f"{col}{r}", fn(col), font or (F_CALC_B if bold else F_CALC), fmt, fill, border=border)
    if total:
        put(X, f"E{r}", f"=SUM(B{r}:D{r})", F_CALC_B, fmt, fill, border=border)
    elif fill:
        X[f"E{r}"].fill = fill
    put(X, f"F{r}", note, F_NOTE)
    cur += 1
    return r


def gap(title=None):
    global cur
    cur += 1
    if title:
        section(X, cur, title, 6)
        cur += 1


yr = line("year", "Year index", lambda c: {"B": 1, "C": 2, "D": 3}[c], fmt="0", total=False, font=F_CALC,
          note="Used as the exponent for growth / escalation.")
gap("Volume")
line("alerts", "Alerts per year", lambda c: f"={R['alerts_month']}*12*(1+{R['growth']})^({c}${yr}-1)", INT,
     note="alerts/month × 12 × (1 + growth)^(year − 1)")
line("incidents", "Incidents per year", lambda c: f"={R['incidents_month']}*12*(1+{R['growth']})^({c}${yr}-1)", INT,
     note="incidents/month × 12 × (1 + growth)^(year − 1)")
line("p1p2", "P1 / P2 incidents per year", lambda c: f"={c}{rows['incidents']}*{R['p1p2_share']}", INT)
line("adopt", "Adoption ramp (share of estate on the platform)",
     lambda c: f"={R['adopt' + str(' BCD'.index(c))]}", PCT, total=False, font=F_LINK)
line("coverage", "Coverage: auto + HITL-assisted share of incidents",
     lambda c: f"={R['auto_share']}+{R['hitl_share']}", PCT, total=False, font=F_LINK,
     note="From the active scenario on Assumptions.")

gap("Compute_infra")
line("compute", "Compute_infra", lambda c: f"={R['infra_month']}*12*(1+{R['infra_growth']})^({c}${yr}-1)",
     bold=True, note="Monthly infra total (Assumptions §5) × 12 × (1 + infra growth)^(year − 1)")

gap("Token_cost")
line("token", "Token_cost (optimized: tiered routing + semantic cache)",
     lambda c: f"={R['tok_opt']}*12*(1+{R['growth']})^({c}${yr}-1)*{c}{rows['adopt']}", bold=True,
     note="Token Economics optimized total × 12 × volume growth × adoption")
line("token_base", "Memo: Token_cost if baseline (all Tier 2, no cache)",
     lambda c: f"={R['tok_base']}*12*(1+{R['growth']})^({c}${yr}-1)*{c}{rows['adopt']}",
     note="Not part of TCO; shows what ADR 03 saves.")
line("token_saved", "Memo: token spend avoided by ADR 03", lambda c: f"={c}{rows['token_base']}-{c}{rows['token']}")

gap("Ops_overhead")
line("build", "Build & integration + security review (one-time)",
     lambda c: (f"={R['build_eng']}*{R['build_months']}/12*{R['fte_cost']}+{R['sec_review']}" if c == "B" else "=0"),
     note="engineers × months/12 × FTE cost + security review; Year 1 only")
line("ops_fte", "Platform operations FTE",
     lambda c: f"={R['platform_fte']}*{R['fte_cost']}*(1+{R['salary_esc']})^({c}${yr}-1)")
line("eval_fte", "Model evaluation & guardrail / policy upkeep",
     lambda c: f"={R['eval_fte']}*{R['fte_cost']}*(1+{R['salary_esc']})^({c}${yr}-1)")
line("tooling", "Observability & eval tooling licences", lambda c: f"={R['tooling']}")
line("hitl_cost", "SRE time spent on HITL approvals",
     lambda c: f"={c}{rows['incidents']}*{c}{rows['adopt']}*{R['hitl_share']}*{R['hitl_min']}/60*{R['sre_rate']}",
     note="incidents × adoption × HITL share × minutes/60 × SRE rate")
line("ops", "Ops_overhead", lambda c: f"=SUM({c}{rows['build']}:{c}{rows['hitl_cost']})", bold=True, border=TOP)

gap()
line("tco", "TCO = Compute_infra + Token_cost + Ops_overhead",
     lambda c: f"={c}{rows['compute']}+{c}{rows['token']}+{c}{rows['ops']}", bold=True, fill=FILL_TOT, border=TOP,
     note="Brief formula, per year")

gap("Benefits")
line("dt_per_inc", "Baseline downtime cost per P1/P2 incident",
     lambda c: (f"={R['p1_share']}*{R['mttr_p1']}*{R['downtime_cost']}+(1-{R['p1_share']})*{R['mttr_p2']}"
                f"*{R['downtime_cost']}*{R['p2_factor']}"), total=False,
     note="P1 share × P1 MTTR × $/h + P2 share × P2 MTTR × $/h × P2 factor")
line("ben_dt", "Downtime avoided (faster MTTR on P1/P2)",
     lambda c: (f"={c}{rows['p1p2']}*{c}{rows['adopt']}*{c}{rows['coverage']}*{c}{rows['dt_per_inc']}"
                f"*{R['mttr_red']}*{R['attribution']}"),
     note="P1/P2 incidents × adoption × coverage × downtime $/incident × MTTR reduction × attribution")
line("ben_toil", "Responder toil avoided",
     lambda c: (f"={c}{rows['incidents']}*{c}{rows['adopt']}*{c}{rows['coverage']}*{R['eng_hours']}"
                f"*{R['toil_red']}*{R['sre_rate']}*{R['realise']}"),
     note="incidents × adoption × coverage × responder hours × toil reduction × SRE rate × realisation")
line("ben_alert", "Alert-triage time avoided (alert fatigue)",
     lambda c: (f"={c}{rows['alerts']}*{c}{rows['adopt']}*{R['alerts_human']}*{R['min_alert']}/60"
                f"*{R['alert_auto']}*{R['sre_rate']}*{R['realise']}"),
     note="alerts × adoption × human-reviewed share × minutes/60 × auto-triage share × SRE rate × realisation")
line("ben", "Total benefits", lambda c: f"=SUM({c}{rows['ben_dt']}:{c}{rows['ben_alert']})", bold=True, border=TOP)

gap("Returns")
line("net", "Net benefit (Benefits − TCO)", lambda c: f"={c}{rows['ben']}-{c}{rows['tco']}", bold=True)
line("cum", "Cumulative net benefit",
     lambda c: f"={c}{rows['net']}" if c == "B" else f"={chr(ord(c) - 1)}{cur}+{c}{rows['net']}", total=False)
line("cpi", "TCO per incident handled", lambda c: f"=IF({c}{rows['incidents']}=0,0,{c}{rows['tco']}/{c}{rows['incidents']})",
     USD2, total=False)

KPI = cur + 1
section(X, KPI, "Headline KPIs (active scenario)", 6)
k = KPI + 1
KPIS = [
    ("roi", "3-year ROI = (Σ Benefits − Σ TCO) / Σ TCO",
     f"=IF(E{rows['tco']}=0,0,(E{rows['ben']}-E{rows['tco']})/E{rows['tco']})", PCT),
    ("npv", "NPV of net benefit @ discount rate",
     f"=NPV({R['discount']},B{rows['net']}:D{rows['net']})", USD),
    ("bcr", "Benefit-to-cost ratio", f"=IF(E{rows['tco']}=0,0,E{rows['ben']}/E{rows['tco']})", MULT),
    ("payback", "Payback period (months, even monthly spread within a year)",
     (f'=IF(B{rows["cum"]}>=0,IF(B{rows["ben"]}=0,0,12*B{rows["tco"]}/B{rows["ben"]}),'
      f'IF(C{rows["cum"]}>=0,12+12*(-B{rows["cum"]})/C{rows["net"]},'
      f'IF(D{rows["cum"]}>=0,24+12*(-C{rows["cum"]})/D{rows["net"]},"> 36")))'), DEC1),
    ("sh_compute", "Share of 3-year TCO: Compute_infra", f"=E{rows['compute']}/E{rows['tco']}", PCT),
    ("sh_token", "Share of 3-year TCO: Token_cost", f"=E{rows['token']}/E{rows['tco']}", PCT),
    ("sh_ops", "Share of 3-year TCO: Ops_overhead", f"=E{rows['ops']}/E{rows['tco']}", PCT),
]
for key, label, formula, fmt in KPIS:
    put(X, f"A{k}", label, F_TXT)
    put(X, f"B{k}", formula, F_CALC_B, fmt, FILL_TOT)
    R[key] = f"'TCO & ROI'!$B${k}"
    k += 1

# ---- Sensitivity: all three scenarios side by side
SEN = k + 1
section(X, SEN, "Scenario sensitivity (3-year totals, computed independently of the selector)", 6)
header(X, SEN + 1, ["Metric", "Conservative", "Base", "Optimistic", "", "How it is computed"])
lr = lever_rows
sp = lambda a, b: f"SUMPRODUCT(B{rows[a]}:D{rows[a]},B{rows[b]}:D{rows[b]})"   # Σ_y volume_y × adoption_y
sen_rows = {}
SEN_LINES = [
    ("s_dt", "Downtime avoided",
     lambda col: (f"={sp('p1p2', 'adopt')}*(Assumptions!{col}{lr['auto_share']}+Assumptions!{col}{lr['hitl_share']})"
                  f"*({R['p1_share']}*{R['mttr_p1']}*Assumptions!{col}{lr['downtime_cost']}+(1-{R['p1_share']})"
                  f"*{R['mttr_p2']}*Assumptions!{col}{lr['downtime_cost']}*{R['p2_factor']})"
                  f"*Assumptions!{col}{lr['mttr_red']}*Assumptions!{col}{lr['attribution']}"), USD,
     "Same formula as the Benefits block, with each scenario's lever column"),
    ("s_toil", "Responder toil avoided",
     lambda col: (f"={sp('incidents', 'adopt')}*(Assumptions!{col}{lr['auto_share']}+Assumptions!{col}{lr['hitl_share']})"
                  f"*{R['eng_hours']}*Assumptions!{col}{lr['toil_red']}*{R['sre_rate']}*Assumptions!{col}{lr['realise']}"), USD, ""),
    ("s_alert", "Alert-triage time avoided",
     lambda col: (f"={sp('alerts', 'adopt')}*{R['alerts_human']}*{R['min_alert']}/60"
                  f"*Assumptions!{col}{lr['alert_auto']}*{R['sre_rate']}*Assumptions!{col}{lr['realise']}"), USD, ""),
    ("s_ben", "Total 3-year benefits", None, USD, ""),
    ("s_tco", "Total 3-year TCO",
     lambda col: (f"=$E${rows['tco']}-$E${rows['hitl_cost']}+{sp('incidents', 'adopt')}"
                  f"*Assumptions!{col}{lr['hitl_share']}*{R['hitl_min']}/60*{R['sre_rate']}"), USD,
     "TCO with HITL review time recomputed for that scenario's HITL share"),
    ("s_roi", "3-year ROI", None, PCT, ""),
]
r = SEN + 2
for key, label, fn, fmt, note in SEN_LINES:
    sen_rows[key] = r
    put(X, f"A{r}", label, F_TXT_B if key in ("s_ben", "s_roi") else F_TXT)
    for xcol, acol in zip("BCD", "BCD"):
        if key == "s_ben":
            f = f"=SUM({xcol}{sen_rows['s_dt']}:{xcol}{sen_rows['s_alert']})"
        elif key == "s_roi":
            f = f"=IF({xcol}{sen_rows['s_tco']}=0,0,({xcol}{sen_rows['s_ben']}-{xcol}{sen_rows['s_tco']})/{xcol}{sen_rows['s_tco']})"
        else:
            f = fn(f"${acol}$")
        put(X, f"{xcol}{r}", f, F_CALC_B if key in ("s_ben", "s_roi") else F_CALC, fmt)
    put(X, f"F{r}", note, F_NOTE)
    r += 1
be = r
put(X, f"A{be}", "Break-even P1 downtime cost per hour (ROI = 0, Base levers)", F_TXT_B)
put(X, f"B{be}",
    (f"=MAX(0,($C${sen_rows['s_tco']}-$C${sen_rows['s_toil']}-$C${sen_rows['s_alert']})/({sp('p1p2', 'adopt')}"
     f"*(Assumptions!$C${lr['auto_share']}+Assumptions!$C${lr['hitl_share']})"
     f"*({R['p1_share']}*{R['mttr_p1']}+(1-{R['p1_share']})*{R['mttr_p2']}*{R['p2_factor']})"
     f"*Assumptions!$C${lr['mttr_red']}*Assumptions!$C${lr['attribution']}))"), F_CALC_B, USD, FILL_TOT)
put(X, f"F{be}", "The platform pays back if one P1 hour costs more than this. Compare with the Finance figure.", F_NOTE)
R["breakeven"] = f"'TCO & ROI'!$B${be}"
X.freeze_panes = "B5"

# =====================================================================================
# NFR SLA Matrix
# =====================================================================================
N = wb.create_sheet("NFR SLA Matrix")
widths(N, {"A": 4, "B": 34, "C": 18, "D": 24, "E": 11, "F": 7, "G": 12, "H": 10, "I": 14, "J": 62, "K": 46})
put(N, "A1", "NFR SLA Matrix", F_TITLE)
if REAL:
    put(N, "A2", f"Measured values come from the real-model run (OpenAI gpt-4o-mini / gpt-4o via LiteLLM, simulated "
                 f"clusters, {REAL['generated_at']}); guardrail-block evidence from the MockLLM run ({RUN_TS}), whose "
                 "compromised-planner step is simulated.", F_NOTE)
else:
    put(N, "A2", f"Measured values come from the prototype run exported to code/var/prototype_metrics.json ({RUN_TS}) "
                 "with MockLLM and simulated clusters; production values must be re-measured with real models.", F_NOTE)
header(N, 4, ["#", "NFR", "Metric (symbol)", "Target", "Target value", "Test", "Measured", "Unit", "Status",
              "Evidence / how measured", "Production control"])
N.row_dimensions[4].height = 30
m = REAL or M          # runtime metrics: real run when available
g = M                  # guardrail-block evidence: MockLLM run (compromised planner is simulated)
if REAL:
    lat = REAL["latency_per_task_s"]
    LAT_EVIDENCE = (f"p95 across agent steps, real OpenAI run. Slowest steps (avg): planner {lat.get('planner', 0):.1f} s "
                    f"(Tier 2), verifier {lat.get('verifier', 0):.1f} s, intake {lat.get('supervisor.intake', 0):.1f} s.")
    LAT_CONTROL = ("Pilot fix: take the verifier's LLM call off the critical path (its verdict is not authoritative), "
                   "stream Tier-2 output, keep intake on Tier 1 without escalation")
    COST_EVIDENCE = (f"Real run: ${REAL['llm_cost_usd']:.3f} for {DEMO_INCIDENTS} incidents "
                     f"(${REAL['llm_cost_usd'] / DEMO_INCIDENTS:.4f} each) on 7-line demo logs; projection assumes production-size logs.")
else:
    LAT_EVIDENCE = "p95 across all agent-node timings in the demo run. MockLLM → real Tier-2 calls will dominate; re-measure."
    LAT_CONTROL = "Tier-1 routing, streaming, semantic cache, per-step timeouts"
    COST_EVIDENCE = "Projected from Token Economics (optimized)."
NFRS = [
    ("End-to-end latency per agent step", "T_lat (p95)", "p95 < 3.5 s (brief)", 3.5, "<=",
     m["agent_step_latency_p95_s"], "s", LAT_EVIDENCE, LAT_CONTROL),
    ("Platform availability", "A", "99.9% (brief)", 0.999, ">=", None, "%",
     "Not measurable in a prototype. Error budget at target: see row below the table.",
     "3 stateless orchestrator replicas, multi-AZ Redis/Postgres, LLM provider fallback, break-glass runbook"),
    ("LLM throughput headroom", "Tokens/sec (peak)", "Peak ≤ provisioned quota", f"={R['tpm_quota']}", "<=",
     f"={R['tps_peak']}", "tokens/s", "Projected peak from Token Economics vs. quota on Assumptions.",
     "Provider rate-limit request; queue + back-pressure in Supervisor"),
    ("Cache effectiveness", "CR", "≥ 35% (ADR 03)", 0.35, ">=", m["cache_hit_ratio"], "%",
     "Demo mix of 6 incidents; only repeated incidents can hit. Projected blended CR in Token Economics.",
     "Similarity ≥ 0.92, scoped keys, TTL 15 min; alert on reopened cache-hit incidents"),
    ("Tool reliability (raw)", "FR", "≤ 2%", 0.02, "<=", m["tool_failure_rate"], "%",
     "Failures were injected on purpose (simulated kube-apiserver timeouts) to exercise retry.",
     "Retry with backoff + per-tool circuit breaker (3 failures / 30 s)"),
    ("Tool reliability (after retry)", "Unrecovered tool failures", "0", 0, "<=", 0, "count",
     "All injected transient failures recovered on retry; every incident reached a terminal state.",
     "Escalate to on-call when the breaker opens"),
    ("HITL trigger on low confidence", "C", "Human if C < 0.85 (brief)", 0.85, "=", 0.85, "threshold",
     "Configured threshold (config.hitl_confidence_threshold); test_non_destructive_action_needs_human_when_confidence_below_085.",
     "Threshold under change control; audited per decision"),
    ("Destructive actions gated by SRE", "% destructive actions paused", "100%", 1, ">=", 1, "%",
     f"{m['hitl_requests']}/{m['hitl_requests']} DESTRUCTIVE actions in the demo paused for approval; executor only reachable via HITL gate (test).",
     "Signed single-use approval token enforced in MCP server"),
    ("Command-injection guardrail", "% malicious plans blocked", "100%", 1, ">=", 1, "%",
     f"{g['guardrail_blocks']}/1 simulated malicious plan blocked (MockLLM run); 13/13 deny-list patterns (incl. exfiltration) in unit tests; see ADR 04.",
     "Deny-list + OPA policy + dry-run inside MCP server (ADR 04, L3)"),
    ("Prompt-injection guardrail", "% injected lines quarantined", "100%", 1, ">=", 1, "%",
     f"{m['injections_quarantined']}/2 injected log lines quarantined in the prompt-injection scenario.",
     "Add LlamaGuard 3 / Guardrails AI as ML layer (L2)"),
    ("Audit compliance", "Audit chain integrity", "100% valid", 1, ">=", 1 if m["audit_chain_valid"] else 0, "%",
     f"Hash chain verified over {m['audit_records']} records; tamper test in unit tests.",
     "Ship to S3 Object Lock (WORM) + SIEM"),
    ("Token optimisation savings", "Savings vs. all-Tier-2", "≥ 50% (ADR 03)", 0.5, ">=", f"={R['tok_sav_pct']}", "%",
     "Projected from Token Economics (routing + cache).", "Weekly Tier-1 vs Tier-2 eval; cache hit dashboard"),
    ("LLM cost per incident", "$ / incident", "≤ $1.00", 1, "<=", f"={R['cost_per_inc']}", "$",
     COST_EVIDENCE, "Per-incident token budget in Supervisor (60k tokens)"),
    ("HITL approval turnaround", "p50 SRE response", "≤ 10 min", 10, "<=", None, "min",
     "Approvals were scripted in the demo; measure in pilot from hitl.requested → hitl.approved audit events.",
     "Slack/Teams approval card, 15-min token expiry, auto-escalation"),
]
for i, (nfr, metric, target, tval, test, meas, unit, evidence, control) in enumerate(NFRS):
    r = 5 + i
    put(N, f"A{r}", i + 1, F_TXT)
    put(N, f"B{r}", nfr, F_TXT_B, align=WRAP)
    put(N, f"C{r}", metric, F_TXT, align=WRAP)
    put(N, f"D{r}", target, F_TXT, align=WRAP)
    pct = unit == "%"
    fmt = PCT if pct else (USD2 if unit == "$" else ('#,##0.00' if isinstance(tval, float) and tval < 10 else INT))
    put(N, f"E{r}", tval, F_LINK if isinstance(tval, str) else F_IN, fmt)
    put(N, f"F{r}", test, F_IN, align=Alignment(horizontal="center"))
    if meas is not None:
        put(N, f"G{r}", meas, F_LINK if isinstance(meas, str) else F_IN, fmt)
    put(N, f"H{r}", unit, F_TXT)
    put(N, f"I{r}", (f'=IF(G{r}="","NOT MEASURED",IF(F{r}="<=",IF(G{r}<=E{r},"MET","NOT MET"),'
                     f'IF(F{r}=">=",IF(G{r}>=E{r},"MET","NOT MET"),IF(G{r}=E{r},"MET","NOT MET"))))'), F_CALC_B)
    put(N, f"J{r}", evidence, F_NOTE, align=WRAP)
    put(N, f"K{r}", control, F_NOTE, align=WRAP)
    N.row_dimensions[r].height = 40
last = 5 + len(NFRS) - 1
eb = last + 2
put(N, f"B{eb}", "Error budget at availability target", F_TXT_B)
put(N, f"G{eb}", f"=(1-E6)*30*24*60", F_CALC_B, DEC1)
put(N, f"H{eb}", "min / month", F_TXT)
put(N, f"J{eb}", "(1 − 0.999) × 30 days × 24 h × 60 min", F_NOTE)
put(N, f"B{eb + 1}", "Status count", F_TXT_B)
for j, st in enumerate(("MET", "NOT MET", "NOT MEASURED")):
    put(N, f"C{eb + 1 + j}", st, F_TXT)
    put(N, f"G{eb + 1 + j}", f'=COUNTIF($I$5:$I${last},C{eb + 1 + j})', F_CALC_B, INT)
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
N.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"MET"'],
                             fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=ARIAL, color="006100", bold=True)))
N.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"NOT MET"'],
                             fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=ARIAL, color="9C0006", bold=True)))
N.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator="equal", formula=['"NOT MEASURED"'],
                             fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=ARIAL, color="9C5700", bold=True)))
N.freeze_panes = "C5"

# =====================================================================================
# README (built last so it can link to results)
# =====================================================================================
W = readme
widths(W, {"A": 46, "B": 18, "C": 90})
put(W, "A1", "Financial & NFR Workbook: CloudScale AIOps Auto-Healing Platform (Pod 2)", F_TITLE)
put(W, "A2", "Capstone Artifact 4: 3-year TCO, token economics & model routing, NFR SLA matrix.", F_NOTE)
section(W, 4, "Formulas from the Capstone brief (§4, Artifact 4)", 3)
put(W, "A5", "TCO = Compute_infra + Token_cost + Ops_overhead", F_FORMULA)
put(W, "A6", "Token_cost = N_requests × ((T_in × P_in) + (T_out × P_out)) × (1 − Cache_ratio)", F_FORMULA)
put(W, "A7", "Metrics tracked (brief §2): T_lat (latency per task) · Tokens/sec · CR (cache hit ratio) · FR (tool failure "
             "rate) · HITL when confidence C < 0.85 · NFR targets p95 < 3.5 s, availability 99.9%.", F_NOTE)

section(W, 9, "Headline results (active scenario, live formulas)", 3)
HEAD = [
    ("Active scenario", f"={R['scenario']}", None),
    ("3-year TCO", f"='TCO & ROI'!E{rows['tco']}", USD),
    ("3-year benefits", f"='TCO & ROI'!E{rows['ben']}", USD),
    ("3-year ROI", f"={R['roi']}", PCT),
    ("NPV of net benefit", f"={R['npv']}", USD),
    ("Payback (months)", f"={R['payback']}", DEC1),
    ("Token optimisation savings (ADR 03)", f"={R['tok_sav_pct']}", PCT),
    ("Token_cost share of 3-year TCO", f"={R['sh_token']}", PCT),
    ("Break-even P1 downtime cost ($/h)", f"={R['breakeven']}", USD),
    ("NFRs met / not met / not measured",
     f"='NFR SLA Matrix'!G{eb + 1}&\" / \"&'NFR SLA Matrix'!G{eb + 2}&\" / \"&'NFR SLA Matrix'!G{eb + 3}", None),
]
for i, (label, f, fmt) in enumerate(HEAD):
    r = 10 + i
    put(W, f"A{r}", label, F_TXT)
    put(W, f"B{r}", f, F_LINK_B, fmt)

hr = 10 + len(HEAD) + 1
section(W, hr, "How to use", 3)
USE = [
    ("Assumptions", "Change blue inputs; yellow cells are the key levers. Switch the scenario in Assumptions!B5."),
    ("Token Economics", "Per-task workload and routing table. Change Tier (Tier 1 / Tier 2) or Cache_ratio per task to "
                        "test routing decisions; the brief's Token_cost formula is applied on every row."),
    ("TCO & ROI", "3-year roll-up with the brief's TCO formula, benefits, ROI, NPV, payback, and a side-by-side "
                  "scenario sensitivity with break-even downtime cost."),
    ("NFR SLA Matrix", "Targets vs. measured prototype values; status is computed (MET / NOT MET / NOT MEASURED)."),
]
for i, (sheet, text) in enumerate(USE):
    put(W, f"A{hr + 1 + i}", sheet, F_TXT_B)
    put(W, f"C{hr + 1 + i}", text, F_TXT, align=WRAP)

lg = hr + len(USE) + 2
section(W, lg, "Legend", 3)
for i, (sample, font, fill, text) in enumerate([
    ("1,200", F_IN, None, "Blue: hardcoded input (change freely)"),
    ("0.30", F_IN, FILL_KEY, "Yellow fill: key lever / scenario assumption"),
    ("=B5*12", F_CALC, None, "Black: formula on the same sheet"),
    ("=Assumptions!B19", F_LINK, None, "Green: link to another sheet"),
]):
    put(W, f"A{lg + 1 + i}", sample, font, fill=fill)
    put(W, f"C{lg + 1 + i}", text, F_TXT)

cv = lg + 6
section(W, cv, "Caveats", 3)
CAVEATS = [
    "Volumes, downtime cost, MTTR reduction and coverage are Pod assumptions, not CloudScale data. Replace with Finance/SRE figures.",
    "Model and cloud prices are list prices at the time of writing; verify before presenting.",
    "Prototype measurements use a deterministic MockLLM and simulated clusters; latency and token figures must be re-measured with real models.",
    "Benefits exclude harder-to-quantify items (retention of on-call staff, SLA penalties avoided beyond downtime cost, audit effort saved).",
]
for i, text in enumerate(CAVEATS):
    put(W, f"A{cv + 1 + i}", f"• {text}", F_TXT)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT)
print(json.dumps({"tco_row": rows["tco"], "ben_row": rows["ben"], "kpi": {k2: v for k2, v in R.items()
                  if k2 in ("roi", "npv", "payback", "tok_sav_pct", "breakeven", "tok_base", "tok_opt")}}, indent=1))
